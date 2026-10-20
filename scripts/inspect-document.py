import sys, json
from pathlib import Path
p = Path(sys.argv[1])
limit = int(sys.argv[2])
chunks = []
metadata = {}
if p.suffix.lower() == '.pdf':
    from pypdf import PdfReader
    doc = PdfReader(p)
    metadata['pages'] = len(doc.pages)
    for i, page in enumerate(doc.pages):
        if sum(map(len, chunks)) >= limit: break
        chunks.append(f'Página {i+1}\n' + (page.extract_text() or '[Página sem texto extraível; pode exigir OCR.]'))
elif p.suffix.lower() == '.docx':
    from docx import Document
    doc = Document(p)
    chunks += [x.text for x in doc.paragraphs if x.text]
    for table in doc.tables:
        chunks += [' | '.join(cell.text for cell in row.cells) for row in table.rows]
elif p.suffix.lower() == '.xlsx':
    from openpyxl import load_workbook
    doc = load_workbook(p, read_only=True, data_only=False)
    metadata['sheets'] = doc.sheetnames
    for sheet in doc:
        chunks.append(f'Planilha: {sheet.title} ({sheet.max_row} linhas, {sheet.max_column} colunas)')
        for row in sheet.iter_rows(values_only=True):
            chunks.append(' | '.join('' if v is None else str(v) for v in row))
            if sum(map(len, chunks)) >= limit: break
        if sum(map(len, chunks)) >= limit: break
    doc.close()
text = '\n'.join(chunks)
print(json.dumps({'path': str(p), **metadata, 'text': text[:limit], 'truncated': len(text) > limit or (metadata.get('pages', 0) > len(chunks))}, ensure_ascii=False))
